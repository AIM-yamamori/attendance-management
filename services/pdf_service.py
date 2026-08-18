"""
pdf_service.py

【概要】
Excel（原本シート）から PDF/ZIP を生成するモジュール（基本設計書5.2.6節・9章）。

採用方式（案A：LibreOffice経由。基本設計書9.2.1節のリスク対応方針に
基づき、案B（reportlab自前描画）から切り替え）：

reportlabで罫線・列幅・数式結果を自前再現する方式は、Excel側の
色・書式・配置・数式の再現が構造的に難しく、数式に依存する項目
（日付・曜日・自動計算列・合計欄等）が空白になる問題が繰り返し
発生した。そのため、サーバー環境にLibreOfficeを導入し、
「Excelファイルをheadlessモードで実際に開いてPDF変換する」方式に
切り替える。LibreOffice自体がExcelの数式・書式・色・配置を解釈して
描画するため、Python側での個別の再現ロジックが不要になる。

処理の流れ：
1. 対象の月次Excelファイルを取得する（attendance_service経由）。
2. 元ファイルには一切手を加えず、一時コピーに対してのみ
   印刷範囲（J1:AQ50）・用紙サイズ（A4縦）をopenpyxlで設定する。
3. LibreOfficeのheadlessモードでコピーをPDFに変換する
   （soffice --headless --convert-to pdf）。印刷範囲が設定されている
   ため、変換されたPDFは自動的にJ1:AQ50の範囲のみとなる。
4. 生成されたPDFバイナリを読み込み、一時ファイルを削除して返す。

出力範囲：原本シート J1:AQ50（要件定義書4.7節）、A4縦、1ユーザー1ファイル。
複数ユーザー選択時はZIPにまとめる。OneDrive・サーバーいずれにも
永続保存せず、生成のたびに一時ディレクトリを使い捨てる
（要件定義書4.7節・5.1節）。
"""

import io
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.worksheet.page import PageMargins

from adapters import excel_adapter
from services import attendance_service

# ============================================
# 出力範囲・印刷設定（要件定義書4.7節・9.1節）
# ============================================

PDF_OUTPUT_RANGE = "J1:AQ50"

# LibreOfficeの実行コマンド。環境によっては "libreoffice" の場合もあるため
# 両方を試すフォールバックを持たせる（14章「未確定事項」相当：
# 実際のサーバー環境で確定したコマンド名に合わせて調整する）。
_SOFFICE_CANDIDATES = ("soffice", "libreoffice")
_CONVERT_TIMEOUT_SECONDS = 60


class PdfConversionError(Exception):
    """PDF変換処理（LibreOffice呼び出し）に失敗した場合に送出する例外。"""


# ============================================
# 印刷範囲・用紙設定（元ファイルには一切書き込まない）
# ============================================

def _prepare_print_ready_copy(
    source_path: str,
    work_dir: Path,
    attendance_data: attendance_service.AttendanceData,
) -> Path:
    """
    月次Excelファイルの一時コピーを作り、そのコピーに対してのみ
    印刷範囲（J1:AQ50）・用紙サイズ（A4縦）・余白を設定して保存する。

    PDF化の対象は「原本」シートのみとする（要件定義書4.7節・9.1節）。
    月次ファイルには「記入例」シート等も含まれているが、LibreOfficeの
    --convert-to pdf はデフォルトでブック内の全シートを変換対象に
    してしまうため、一時コピーから「原本」シート以外を削除してから
    LibreOfficeに渡すことで、確実に原本シートのみのPDFにする。

    さらに、C列（曜日）はExcel側の数式 =TEXT(B14,"aaa") が
    LibreOfficeのロケール設定に依存して英語表記（Mon等）になって
    しまう問題があるため、attendance_service が既に日本語で
    再計算済みの値（DayEntry.weekday）でC列を直接上書きする。

    元ファイル（OneDrive上のファイルのローカルコピー）は一切変更しない。
    このコピーに対してのみ値・印刷設定・シート構成を書き込む。
    """
    copy_path = work_dir / "print_ready.xlsx"
    shutil.copyfile(source_path, copy_path)

    workbook = openpyxl.load_workbook(copy_path)

    # 「原本」シート以外は削除し、PDF化の対象を原本シートのみにする
    for sheet_name in list(workbook.sheetnames):
        if sheet_name != excel_adapter.SHEET_HONBUN:
            del workbook[sheet_name]

    workbook.active = workbook.sheetnames.index(excel_adapter.SHEET_HONBUN)

    ws = workbook[excel_adapter.SHEET_HONBUN]

    ws.print_area = PDF_OUTPUT_RANGE
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

    if ws.sheet_properties.pageSetUpPr is None:
        from openpyxl.worksheet.properties import PageSetupProperties
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    else:
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.3, bottom=0.3, header=0, footer=0)

    # C列（曜日）を、LibreOfficeのロケールに依存しないよう
    # Python側で計算済みの日本語表記で直接上書きする
    weekday_col_letter = "C"
    for entry in attendance_data.entries:
        if entry.date_value is None:
            continue
        cell = ws[f"{weekday_col_letter}{entry.row}"]
        cell.value = entry.weekday or ""

    workbook.save(copy_path)
    return copy_path


# ============================================
# LibreOffice呼び出し
# ============================================

def _find_soffice_command() -> str:
    """
    LibreOfficeの実行コマンドを探す。soffice / libreoffice のいずれかが
    PATH上に存在すればそれを使う。どちらも見つからない場合は
    PdfConversionErrorを送出する。
    """
    for candidate in _SOFFICE_CANDIDATES:
        if shutil.which(candidate):
            return candidate
    raise PdfConversionError(
        "LibreOffice（soffice/libreoffice）が見つかりません。"
        "サーバー環境にLibreOfficeがインストールされているか確認してください。"
    )


def _convert_to_pdf(xlsx_path: Path, work_dir: Path) -> Path:
    """
    LibreOfficeのheadlessモードで xlsx_path をPDFに変換し、
    生成されたPDFファイルのパスを返す。

    --convert-to pdf は、入力ファイルと同名（拡張子だけ.pdf）の
    ファイルを --outdir に生成する仕様のため、変換後のファイル名は
    xlsx_path のファイル名から機械的に導出する。
    """
    soffice_cmd = _find_soffice_command()

    result = subprocess.run(
        [
            soffice_cmd,
            "--headless",
            "--norestore",
            "--convert-to", "pdf",
            "--outdir", str(work_dir),
            str(xlsx_path),
        ],
        capture_output=True,
        timeout=_CONVERT_TIMEOUT_SECONDS,
    )

    if result.returncode != 0:
        raise PdfConversionError(
            f"LibreOfficeによるPDF変換に失敗しました: "
            f"{result.stderr.decode('utf-8', errors='ignore')}"
        )

    expected_pdf_path = work_dir / (xlsx_path.stem + ".pdf")
    if not expected_pdf_path.exists():
        raise PdfConversionError(
            f"PDF変換後のファイルが見つかりません: {expected_pdf_path}"
        )

    return expected_pdf_path


# ============================================
# 5.2.6節 公開インターフェース
# ============================================

def generate_pdf(employee_id: str, target_month: str) -> bytes:
    """
    単一ユーザーのPDFを生成する（基本設計書5.2.6節・9章）。
    OneDrive・サーバーいずれにも保存せず、バイナリをそのまま返す。
    """
    from services import user_service

    user = user_service.get_user(employee_id)
    if user is None:
        raise ValueError(f"ユーザーが見つかりません: {employee_id}")

    handle = attendance_service.get_or_create_monthly_file(
        employee_id=employee_id,
        target_month=target_month,
        full_name_no_space=user.full_name_no_space,
        full_name_with_space=user.full_name_with_space,
        department=user.department,
    )

    attendance_data = attendance_service.load_attendance(handle)

    with tempfile.TemporaryDirectory(prefix="pdf-export-") as tmp_dir_str:
        work_dir = Path(tmp_dir_str)

        print_ready_path = _prepare_print_ready_copy(handle.local_path, work_dir, attendance_data)
        pdf_path = _convert_to_pdf(print_ready_path, work_dir)

        return pdf_path.read_bytes()


def build_pdf_file_name(employee_id: str, target_month: str) -> str:
    """
    ダウンロードファイル名を組み立てる（要件定義書5.1節：
    Excelと対になるファイル名規則、拡張子のみ.pdf）。
    """
    from services import user_service

    user = user_service.get_user(employee_id)
    full_name_no_space = user.full_name_no_space if user else employee_id
    return attendance_service.build_monthly_file_name(
        target_month, full_name_no_space
    ).replace(".xlsx", ".pdf")


def generate_zip(employee_ids: list[str], target_month: str) -> bytes:
    """
    複数ユーザー分のPDFをZIPにまとめて返す（基本設計書5.2.6節）。
    1ユーザー1PDFファイルとし、結合はしない（要件定義書4.7節）。
    """
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for employee_id in employee_ids:
            pdf_bytes = generate_pdf(employee_id, target_month)
            file_name = build_pdf_file_name(employee_id, target_month)
            zf.writestr(file_name, pdf_bytes)

    buffer.seek(0)
    return buffer.read()